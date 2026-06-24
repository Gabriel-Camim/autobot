from __future__ import annotations

import difflib
import hashlib
import io
import json
import re
import sqlite3
import threading
import time
from datetime import datetime, timezone
from pathlib import Path, PurePosixPath
from typing import Any, Dict, Iterable, List, Optional
from uuid import uuid4

import yaml
from langchain_core.messages import HumanMessage, SystemMessage

from agent import AppError, _build_llm, _response_content_to_text, build_rag_probe, sanitize_answer_text
from config import Settings
from rag_quality import get_knowledge_suggestion, get_rag_feedback, get_rag_trace, triage_rag_feedback


_lock = threading.Lock()

DRAFT_ROOT = "knowledge/_drafts"
MAX_ATTACHMENTS_PER_DRAFT = 5
MAX_ATTACHMENT_MB = 10
ALLOWED_ATTACHMENT_SUFFIXES = {".md", ".txt", ".json", ".pdf", ".docx"}
CASE_RESOLVED_STATUSES = {"resolved", "archived", "deleted"}
DOC_DONE_STATUSES = {"merged", "ignored"}


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _today() -> str:
    return datetime.now(timezone.utc).date().isoformat()


def _uses_postgres(settings: Settings) -> bool:
    return settings.database_url.strip().startswith(("postgres://", "postgresql://"))


def _placeholder(settings: Settings) -> str:
    return "%s" if _uses_postgres(settings) else "?"


def _safe_json(data: Any) -> str:
    return json.dumps(data if data is not None else {}, ensure_ascii=False, separators=(",", ":"), default=str)


def _load_json(value: Any, fallback: Any = None) -> Any:
    if value in (None, ""):
        return {} if fallback is None else fallback
    try:
        return json.loads(value)
    except (TypeError, json.JSONDecodeError):
        return {} if fallback is None else fallback


def _connect(settings: Settings):
    if _uses_postgres(settings):
        import psycopg
        from psycopg.rows import dict_row

        conn = psycopg.connect(settings.database_url.strip(), row_factory=dict_row, autocommit=True)
        _ensure_schema(conn, postgres=True)
        return conn

    db_path = settings.resolved_events_db_path.parent / "draft_studio.sqlite3"
    db_path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(db_path, timeout=10)
    conn.row_factory = sqlite3.Row
    _ensure_schema(conn, postgres=False)
    return conn


def _ensure_schema(conn, *, postgres: bool) -> None:
    timestamp = "TIMESTAMPTZ" if postgres else "TEXT"
    conn.execute(
        f"""
        CREATE TABLE IF NOT EXISTS curation_cases (
            id TEXT PRIMARY KEY,
            created_at {timestamp} NOT NULL,
            updated_at {timestamp} NOT NULL,
            status TEXT NOT NULL,
            title TEXT NOT NULL,
            instruction TEXT,
            source_kind TEXT NOT NULL,
            source_feedback_id TEXT,
            source_trace_id TEXT,
            validation_json TEXT NOT NULL,
            payload_json TEXT NOT NULL
        )
        """
    )
    conn.execute(
        f"""
        CREATE TABLE IF NOT EXISTS knowledge_drafts (
            id TEXT PRIMARY KEY,
            created_at {timestamp} NOT NULL,
            updated_at {timestamp} NOT NULL,
            status TEXT NOT NULL,
            title TEXT NOT NULL,
            instruction TEXT,
            source_suggestion_id TEXT,
            source_trace_id TEXT,
            suggested_path TEXT NOT NULL,
            git_path TEXT,
            git_commit_sha TEXT,
            canonical_targets_json TEXT NOT NULL,
            draft_markdown TEXT NOT NULL,
            proposed_markdown TEXT,
            evidence_json TEXT NOT NULL,
            conflict_report_json TEXT NOT NULL,
            eval_suggestions_json TEXT NOT NULL,
            payload_json TEXT NOT NULL
        )
        """
    )
    _add_column_if_missing(conn, "knowledge_drafts", "case_id", "TEXT", postgres=postgres)
    _add_column_if_missing(conn, "knowledge_drafts", "source_feedback_id", "TEXT", postgres=postgres)
    _add_column_if_missing(conn, "knowledge_drafts", "target_path", "TEXT", postgres=postgres)
    _add_column_if_missing(conn, "knowledge_drafts", "decision", "TEXT", postgres=postgres)
    _add_column_if_missing(conn, "knowledge_drafts", "completed_at", timestamp, postgres=postgres)
    _add_column_if_missing(conn, "knowledge_drafts", "validation_json", "TEXT", postgres=postgres)
    conn.execute(
        f"""
        CREATE TABLE IF NOT EXISTS draft_runs (
            id TEXT PRIMARY KEY,
            draft_id TEXT NOT NULL,
            created_at {timestamp} NOT NULL,
            model TEXT,
            instruction TEXT,
            status TEXT NOT NULL,
            input_json TEXT NOT NULL,
            output_json TEXT NOT NULL,
            metrics_json TEXT NOT NULL,
            error TEXT
        )
        """
    )
    conn.execute(
        f"""
        CREATE TABLE IF NOT EXISTS draft_attachments (
            id TEXT PRIMARY KEY,
            draft_id TEXT NOT NULL,
            created_at {timestamp} NOT NULL,
            filename TEXT NOT NULL,
            content_type TEXT,
            size_bytes INTEGER NOT NULL,
            sha256 TEXT NOT NULL,
            extracted_text TEXT NOT NULL,
            metadata_json TEXT NOT NULL
        )
        """
    )
    conn.execute(
        f"""
        CREATE TABLE IF NOT EXISTS draft_patches (
            id TEXT PRIMARY KEY,
            draft_id TEXT NOT NULL,
            created_at {timestamp} NOT NULL,
            status TEXT NOT NULL,
            target_path TEXT NOT NULL,
            original_content TEXT NOT NULL,
            proposed_content TEXT NOT NULL,
            diff_text TEXT NOT NULL,
            commit_sha TEXT,
            payload_json TEXT NOT NULL
        )
        """
    )
    for statement in (
        "CREATE INDEX IF NOT EXISTS idx_curation_cases_status ON curation_cases(status)",
        "CREATE INDEX IF NOT EXISTS idx_curation_cases_feedback ON curation_cases(source_feedback_id)",
        "CREATE INDEX IF NOT EXISTS idx_curation_cases_updated_at ON curation_cases(updated_at)",
        "CREATE INDEX IF NOT EXISTS idx_knowledge_drafts_status ON knowledge_drafts(status)",
        "CREATE INDEX IF NOT EXISTS idx_knowledge_drafts_created_at ON knowledge_drafts(created_at)",
        "CREATE INDEX IF NOT EXISTS idx_knowledge_drafts_suggestion ON knowledge_drafts(source_suggestion_id)",
        "CREATE INDEX IF NOT EXISTS idx_knowledge_drafts_case ON knowledge_drafts(case_id)",
        "CREATE INDEX IF NOT EXISTS idx_knowledge_drafts_feedback ON knowledge_drafts(source_feedback_id)",
        "CREATE INDEX IF NOT EXISTS idx_draft_runs_draft_id ON draft_runs(draft_id)",
        "CREATE INDEX IF NOT EXISTS idx_draft_attachments_draft_id ON draft_attachments(draft_id)",
        "CREATE INDEX IF NOT EXISTS idx_draft_patches_draft_id ON draft_patches(draft_id)",
    ):
        conn.execute(statement)
    _backfill_curation_cases(conn, postgres=postgres)


def _table_columns(conn, table: str, *, postgres: bool) -> set[str]:
    if postgres:
        rows = conn.execute(
            """
            SELECT column_name FROM information_schema.columns
            WHERE table_name = %s AND table_schema = current_schema()
            """,
            (table,),
        ).fetchall()
        return {row["column_name"] if isinstance(row, dict) else row[0] for row in rows}
    rows = conn.execute(f"PRAGMA table_info({table})").fetchall()
    return {row["name"] if isinstance(row, dict) else row[1] for row in rows}


def _add_column_if_missing(conn, table: str, column: str, definition: str, *, postgres: bool) -> None:
    if column in _table_columns(conn, table, postgres=postgres):
        return
    conn.execute(f"ALTER TABLE {table} ADD COLUMN {column} {definition}")


def _backfill_curation_cases(conn, *, postgres: bool) -> None:
    marker = "%s" if postgres else "?"
    rows = conn.execute("SELECT * FROM knowledge_drafts WHERE case_id IS NULL OR case_id = ''").fetchall()
    for row in rows:
        if not isinstance(row, dict):
            row = dict(row)
        case_id = str(uuid4())
        now = _now()
        title = row.get("title") or "Caso de curadoria"
        source_kind = "suggestion" if row.get("source_suggestion_id") else "manual"
        source_feedback_id = row.get("source_feedback_id")
        source_trace_id = row.get("source_trace_id")
        validation = {"backfilled": True}
        payload = {"backfilled_from_draft": row.get("id")}
        conn.execute(
            f"""
            INSERT INTO curation_cases
            (id, created_at, updated_at, status, title, instruction, source_kind, source_feedback_id,
             source_trace_id, validation_json, payload_json)
            VALUES ({",".join([marker] * 11)})
            """,
            (
                case_id,
                now,
                now,
                row.get("status") or "review",
                title,
                row.get("instruction"),
                source_kind,
                source_feedback_id,
                source_trace_id,
                _safe_json(validation),
                _safe_json(payload),
            ),
        )
        conn.execute(
            f"UPDATE knowledge_drafts SET case_id = {marker}, decision = COALESCE(decision, {marker}), validation_json = COALESCE(validation_json, {marker}) WHERE id = {marker}",
            (case_id, "edit", _safe_json({}), row.get("id")),
        )


def init_draft_studio_db(settings: Settings) -> None:
    with _lock:
        with _connect(settings):
            pass


def draft_studio_storage_status(settings: Settings) -> Dict[str, Any]:
    try:
        init_draft_studio_db(settings)
        return {"backend": "postgres" if _uses_postgres(settings) else "sqlite", "ok": True, "error": None}
    except Exception as exc:
        return {"backend": "postgres" if _uses_postgres(settings) else "sqlite", "ok": False, "error": str(exc)[:240]}


def _slugify(text: str, fallback: str = "draft") -> str:
    normalized = re.sub(r"[^a-zA-Z0-9]+", "-", (text or "").lower()).strip("-")
    return (normalized or fallback)[:72].strip("-") or fallback


def _normalize_repo_path(raw_path: str) -> str:
    cleaned = (raw_path or "").strip().replace("\\", "/").strip("/")
    if not cleaned:
        raise AppError("invalid_path", "Informe um caminho valido.", 400)
    path = PurePosixPath(cleaned)
    if path.is_absolute() or any(part in {"", ".", ".."} for part in path.parts):
        raise AppError("invalid_path", "Caminho invalido.", 400)
    return path.as_posix()


def _draft_path(title: str) -> str:
    return f"{DRAFT_ROOT}/{_today()}-{_slugify(title)}.md"


def _canonical_repo_path_from_source(source: str) -> Optional[str]:
    source = (source or "").strip().replace("\\", "/").strip("/")
    if not source or source.startswith("_drafts/") or source.startswith("reports/"):
        return None
    if source.startswith("knowledge/"):
        repo_path = source
    else:
        repo_path = f"knowledge/{source}"
    if repo_path.startswith(f"{DRAFT_ROOT}/"):
        return None
    return repo_path if repo_path.endswith(".md") else None


def _local_path_for_repo_path(settings: Settings, repo_path: str) -> Path:
    normalized = _normalize_repo_path(repo_path)
    parts = PurePosixPath(normalized).parts
    if not parts or parts[0] != "knowledge":
        raise AppError("invalid_path", "Draft Studio so pode operar em knowledge/.", 400)
    target = (settings.resolved_knowledge_dir / Path(PurePosixPath(*parts[1:]).as_posix())).resolve()
    base = settings.resolved_knowledge_dir.resolve()
    if target != base and base not in target.parents:
        raise AppError("invalid_path", "Caminho fora da base knowledge/.", 400)
    return target


def _frontmatter_for_draft(title: str, summary: str) -> Dict[str, Any]:
    return {
        "title": title,
        "category": "curadoria",
        "tags": ["draft", "ragops", "curadoria"],
        "visibility": "draft",
        "priority": 3,
        "updated_at": _today(),
        "summary": summary,
    }


def _markdown_from_frontmatter(frontmatter: Dict[str, Any], body: str) -> str:
    return "---\n" + yaml.safe_dump(frontmatter, allow_unicode=True, sort_keys=False).strip() + "\n---\n\n" + body.strip() + "\n"


def _initial_markdown(title: str, instruction: str, suggestion: Optional[Dict[str, Any]] = None) -> str:
    if suggestion and suggestion.get("draft_markdown"):
        return str(suggestion["draft_markdown"])
    frontmatter = _frontmatter_for_draft(title, "Rascunho privado para curadoria assistida por IA.")
    body = f"""# {title}

## Instrução original

{instruction or "Sem instrução informada."}

## Rascunho

Este draft ainda não foi refinado pelo agente.
"""
    return _markdown_from_frontmatter(frontmatter, body)


def _draft_from_row(row: Any) -> Optional[Dict[str, Any]]:
    if not row:
        return None
    if not isinstance(row, dict):
        row = dict(row)
    return {
        "id": row.get("id"),
        "created_at": str(row.get("created_at")),
        "updated_at": str(row.get("updated_at")),
        "status": row.get("status"),
        "title": row.get("title"),
        "instruction": row.get("instruction"),
        "source_suggestion_id": row.get("source_suggestion_id"),
        "source_feedback_id": row.get("source_feedback_id"),
        "source_trace_id": row.get("source_trace_id"),
        "case_id": row.get("case_id"),
        "suggested_path": row.get("suggested_path"),
        "git_path": row.get("git_path"),
        "git_commit_sha": row.get("git_commit_sha"),
        "target_path": row.get("target_path"),
        "decision": row.get("decision") or "edit",
        "completed_at": str(row.get("completed_at")) if row.get("completed_at") else None,
        "canonical_targets": _load_json(row.get("canonical_targets_json"), []),
        "draft_markdown": row.get("draft_markdown"),
        "proposed_markdown": row.get("proposed_markdown"),
        "evidence": _load_json(row.get("evidence_json"), []),
        "conflict_report": _load_json(row.get("conflict_report_json"), []),
        "eval_suggestions": _load_json(row.get("eval_suggestions_json"), []),
        "validation": _load_json(row.get("validation_json"), {}),
        "payload": _load_json(row.get("payload_json"), {}),
    }


def _case_from_row(row: Any, drafts: Optional[List[Dict[str, Any]]] = None) -> Optional[Dict[str, Any]]:
    if not row:
        return None
    if not isinstance(row, dict):
        row = dict(row)
    case = {
        "id": row.get("id"),
        "created_at": str(row.get("created_at")),
        "updated_at": str(row.get("updated_at")),
        "status": row.get("status"),
        "title": row.get("title"),
        "instruction": row.get("instruction"),
        "source_kind": row.get("source_kind"),
        "source_feedback_id": row.get("source_feedback_id"),
        "source_trace_id": row.get("source_trace_id"),
        "validation": _load_json(row.get("validation_json"), {}),
        "payload": _load_json(row.get("payload_json"), {}),
    }
    if drafts is not None:
        case["drafts"] = drafts
        case["timeline"] = _case_timeline(case, drafts)
        case["next_action"] = _case_next_action(case, drafts)
    return case


def _case_timeline(case: Dict[str, Any], drafts: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    validation = case.get("validation") or {}
    has_targets = bool(drafts)
    has_ai = any((draft.get("status") in {"ai_generated", "draft_committed", "patch_proposed", "merged"}) for draft in drafts)
    has_draft_git = any(draft.get("git_path") for draft in drafts)
    has_patch = any((draft.get("patches") or []) for draft in drafts)
    docs_done = has_targets and all(draft.get("status") in DOC_DONE_STATUSES for draft in drafts)
    reindexed = (validation.get("reindex") or {}).get("state") == "success"
    eval_state = (validation.get("eval") or {}).get("state")
    steps = [
        ("signal", "Sinal recebido", True),
        ("targets", "Documentos definidos", has_targets),
        ("draft", "IA gerou rascunho", has_ai),
        ("draft_git", "Draft privado salvo", has_draft_git),
        ("patch", "Diff criado", has_patch),
        ("merge", "Documentos concluídos", docs_done),
        ("reindex", "RAG reindexado", reindexed),
        ("eval", "Eval registrado", bool(eval_state)),
    ]
    return [{"id": key, "label": label, "done": bool(done)} for key, label, done in steps]


def _case_next_action(case: Dict[str, Any], drafts: List[Dict[str, Any]]) -> str:
    validation = case.get("validation") or {}
    if not drafts:
        return "Selecione os Markdown que precisam de curadoria."
    pending = [draft for draft in drafts if draft.get("status") not in DOC_DONE_STATUSES]
    if pending:
        return "Revise cada documento: gere com IA, salve o draft, crie o diff e aplique ou ignore."
    if (validation.get("reindex") or {}).get("state") != "success":
        return "Reindexe o RAG para publicar a mudança no índice."
    if case.get("status") != "resolved":
        return "Caso validado por reindex. Você já pode resolver o feedback."
    return "Caso resolvido."


def _attachment_from_row(row: Any) -> Dict[str, Any]:
    if not isinstance(row, dict):
        row = dict(row)
    return {
        "id": row.get("id"),
        "draft_id": row.get("draft_id"),
        "created_at": str(row.get("created_at")),
        "filename": row.get("filename"),
        "content_type": row.get("content_type"),
        "size_bytes": row.get("size_bytes"),
        "sha256": row.get("sha256"),
        "extracted_text": row.get("extracted_text"),
        "metadata": _load_json(row.get("metadata_json"), {}),
    }


def _patch_from_row(row: Any) -> Dict[str, Any]:
    if not isinstance(row, dict):
        row = dict(row)
    return {
        "id": row.get("id"),
        "draft_id": row.get("draft_id"),
        "created_at": str(row.get("created_at")),
        "status": row.get("status"),
        "target_path": row.get("target_path"),
        "original_content": row.get("original_content"),
        "proposed_content": row.get("proposed_content"),
        "diff_text": row.get("diff_text"),
        "commit_sha": row.get("commit_sha"),
        "payload": _load_json(row.get("payload_json"), {}),
    }


def _run_from_row(row: Any) -> Dict[str, Any]:
    if not isinstance(row, dict):
        row = dict(row)
    return {
        "id": row.get("id"),
        "draft_id": row.get("draft_id"),
        "created_at": str(row.get("created_at")),
        "model": row.get("model"),
        "instruction": row.get("instruction"),
        "status": row.get("status"),
        "input": _load_json(row.get("input_json"), {}),
        "output": _load_json(row.get("output_json"), {}),
        "metrics": _load_json(row.get("metrics_json"), {}),
        "error": row.get("error"),
    }


def create_draft(settings: Settings, payload: Dict[str, Any]) -> Dict[str, Any]:
    suggestion_id = payload.get("suggestion_id")
    trace_id = payload.get("trace_id")
    suggestion = get_knowledge_suggestion(settings, suggestion_id) if suggestion_id else None
    if suggestion and not trace_id:
        trace_id = ((suggestion.get("payload") or {}).get("trace") or {}).get("id")
    trace = get_rag_trace(settings, trace_id) if trace_id else None
    title = str(payload.get("title") or (suggestion or {}).get("title") or "Draft de curadoria").strip()
    instruction = str(payload.get("instruction") or "").strip()
    if not instruction:
        instruction = str((suggestion or {}).get("rationale") or (trace or {}).get("question") or "Refinar contexto para o RAG.")
    suggested_path = _normalize_repo_path(payload.get("suggested_path") or (suggestion or {}).get("suggested_path") or _draft_path(title))
    if not suggested_path.startswith(f"{DRAFT_ROOT}/"):
        suggested_path = _draft_path(title)
    now = _now()
    draft = {
        "id": str(uuid4()),
        "created_at": now,
        "updated_at": now,
        "status": payload.get("status") or "review",
        "title": title,
        "instruction": instruction,
        "source_suggestion_id": suggestion_id,
        "source_feedback_id": payload.get("source_feedback_id"),
        "source_trace_id": trace_id,
        "case_id": payload.get("case_id"),
        "suggested_path": suggested_path,
        "git_path": None,
        "git_commit_sha": None,
        "target_path": payload.get("target_path"),
        "decision": payload.get("decision") or "edit",
        "completed_at": None,
        "canonical_targets": payload.get("canonical_targets") or [],
        "draft_markdown": payload.get("draft_markdown") or _initial_markdown(title, instruction, suggestion),
        "proposed_markdown": None,
        "evidence": [],
        "conflict_report": [],
        "eval_suggestions": [],
        "validation": {},
        "payload": {"suggestion": suggestion, "trace": trace},
    }
    marker = _placeholder(settings)
    row = (
        draft["id"],
        draft["created_at"],
        draft["updated_at"],
        draft["status"],
        draft["title"],
        draft["instruction"],
        draft["source_suggestion_id"],
        draft["source_feedback_id"],
        draft["source_trace_id"],
        draft["case_id"],
        draft["suggested_path"],
        draft["git_path"],
        draft["git_commit_sha"],
        draft["target_path"],
        draft["decision"],
        draft["completed_at"],
        _safe_json(draft["canonical_targets"]),
        draft["draft_markdown"],
        draft["proposed_markdown"],
        _safe_json(draft["evidence"]),
        _safe_json(draft["conflict_report"]),
        _safe_json(draft["eval_suggestions"]),
        _safe_json(draft["validation"]),
        _safe_json(draft["payload"]),
    )
    with _lock:
        with _connect(settings) as conn:
            conn.execute(
                f"""
                INSERT INTO knowledge_drafts
                (id, created_at, updated_at, status, title, instruction, source_suggestion_id, source_feedback_id,
                 source_trace_id, case_id, suggested_path, git_path, git_commit_sha, target_path, decision, completed_at,
                 canonical_targets_json, draft_markdown, proposed_markdown, evidence_json, conflict_report_json,
                 eval_suggestions_json, validation_json, payload_json)
                VALUES ({",".join([marker] * 24)})
                """,
                row,
            )
    return get_draft(settings, draft["id"]) or draft


def list_drafts(settings: Settings, *, status: Optional[str] = None, limit: int = 100) -> List[Dict[str, Any]]:
    limit = max(1, min(limit, 500))
    marker = _placeholder(settings)
    params: List[Any] = []
    query = "SELECT * FROM knowledge_drafts"
    if status:
        query += f" WHERE status = {marker}"
        params.append(status)
    query += f" ORDER BY updated_at DESC LIMIT {marker}"
    params.append(limit)
    with _lock:
        with _connect(settings) as conn:
            rows = conn.execute(query, params).fetchall()
    return [item for item in (_draft_from_row(row) for row in rows) if item]


def get_draft(settings: Settings, draft_id: str, *, include_related: bool = False) -> Optional[Dict[str, Any]]:
    marker = _placeholder(settings)
    with _lock:
        with _connect(settings) as conn:
            row = conn.execute(f"SELECT * FROM knowledge_drafts WHERE id = {marker}", (draft_id,)).fetchone()
            draft = _draft_from_row(row)
            if not draft or not include_related:
                return draft
            attachments = conn.execute(
                f"SELECT * FROM draft_attachments WHERE draft_id = {marker} ORDER BY created_at DESC",
                (draft_id,),
            ).fetchall()
            patches = conn.execute(
                f"SELECT * FROM draft_patches WHERE draft_id = {marker} ORDER BY created_at DESC",
                (draft_id,),
            ).fetchall()
            runs = conn.execute(
                f"SELECT * FROM draft_runs WHERE draft_id = {marker} ORDER BY created_at DESC LIMIT 20",
                (draft_id,),
            ).fetchall()
    draft["attachments"] = [_attachment_from_row(row) for row in attachments]
    draft["patches"] = [_patch_from_row(row) for row in patches]
    draft["runs"] = [_run_from_row(row) for row in runs]
    return draft


def update_draft(settings: Settings, draft_id: str, updates: Dict[str, Any]) -> Optional[Dict[str, Any]]:
    allowed = {
        "status": "status",
        "title": "title",
        "instruction": "instruction",
        "source_feedback_id": "source_feedback_id",
        "source_trace_id": "source_trace_id",
        "case_id": "case_id",
        "suggested_path": "suggested_path",
        "git_path": "git_path",
        "git_commit_sha": "git_commit_sha",
        "target_path": "target_path",
        "decision": "decision",
        "completed_at": "completed_at",
        "draft_markdown": "draft_markdown",
        "proposed_markdown": "proposed_markdown",
    }
    marker = _placeholder(settings)
    fields = [f"updated_at = {marker}"]
    params: List[Any] = [_now()]
    for key, column in allowed.items():
        if key in updates:
            value = updates[key]
            if key in {"suggested_path", "git_path", "target_path"} and value:
                value = _normalize_repo_path(str(value))
            fields.append(f"{column} = {marker}")
            params.append(value)
    json_fields = {
        "canonical_targets": "canonical_targets_json",
        "evidence": "evidence_json",
        "conflict_report": "conflict_report_json",
        "eval_suggestions": "eval_suggestions_json",
        "validation": "validation_json",
        "payload": "payload_json",
    }
    for key, column in json_fields.items():
        if key in updates:
            fields.append(f"{column} = {marker}")
            params.append(_safe_json(updates[key]))
    params.append(draft_id)
    with _lock:
        with _connect(settings) as conn:
            conn.execute(f"UPDATE knowledge_drafts SET {', '.join(fields)} WHERE id = {marker}", params)
    return get_draft(settings, draft_id, include_related=True)


def _drafts_for_case(settings: Settings, case_id: str, *, include_related: bool = True) -> List[Dict[str, Any]]:
    marker = _placeholder(settings)
    with _lock:
        with _connect(settings) as conn:
            rows = conn.execute(
                f"SELECT * FROM knowledge_drafts WHERE case_id = {marker} ORDER BY created_at ASC",
                (case_id,),
            ).fetchall()
    drafts = [item for item in (_draft_from_row(row) for row in rows) if item]
    if not include_related:
        return drafts
    return [get_draft(settings, draft["id"], include_related=True) or draft for draft in drafts]


def _case_payload(settings: Settings, case_id: str, *, include_drafts: bool = True) -> Optional[Dict[str, Any]]:
    marker = _placeholder(settings)
    with _lock:
        with _connect(settings) as conn:
            row = conn.execute(f"SELECT * FROM curation_cases WHERE id = {marker}", (case_id,)).fetchone()
    if not row:
        return None
    drafts = _drafts_for_case(settings, case_id, include_related=True) if include_drafts else None
    return _case_from_row(row, drafts)


def list_curation_cases(settings: Settings, *, status: Optional[str] = None, limit: int = 100) -> List[Dict[str, Any]]:
    limit = max(1, min(limit, 500))
    marker = _placeholder(settings)
    params: List[Any] = []
    query = "SELECT * FROM curation_cases"
    if status:
        query += f" WHERE status = {marker}"
        params.append(status)
    else:
        query += " WHERE status != " + marker
        params.append("deleted")
    query += f" ORDER BY updated_at DESC LIMIT {marker}"
    params.append(limit)
    with _lock:
        with _connect(settings) as conn:
            rows = conn.execute(query, params).fetchall()
    cases = []
    for row in rows:
        case = _case_from_row(row, None)
        if case:
            drafts = _drafts_for_case(settings, case["id"], include_related=False)
            case["drafts"] = drafts
            case["timeline"] = _case_timeline(case, drafts)
            case["next_action"] = _case_next_action(case, drafts)
            cases.append(case)
    return cases


def get_curation_case(settings: Settings, case_id: str) -> Optional[Dict[str, Any]]:
    return _case_payload(settings, case_id, include_drafts=True)


def active_case_for_feedback(settings: Settings, feedback_id: str) -> Optional[Dict[str, Any]]:
    marker = _placeholder(settings)
    with _lock:
        with _connect(settings) as conn:
            placeholders = ",".join([marker] * len(CASE_RESOLVED_STATUSES))
            row = conn.execute(
                f"""
                SELECT * FROM curation_cases
                WHERE source_feedback_id = {marker} AND status NOT IN ({placeholders})
                ORDER BY updated_at DESC LIMIT 1
                """,
                (feedback_id, *CASE_RESOLVED_STATUSES),
            ).fetchone()
    if not row:
        return None
    case = _case_from_row(row, None)
    return get_curation_case(settings, case["id"]) if case else None


def _update_case(settings: Settings, case_id: str, updates: Dict[str, Any]) -> Optional[Dict[str, Any]]:
    marker = _placeholder(settings)
    fields = [f"updated_at = {marker}"]
    params: List[Any] = [_now()]
    for key, column in {
        "status": "status",
        "title": "title",
        "instruction": "instruction",
        "source_kind": "source_kind",
        "source_feedback_id": "source_feedback_id",
        "source_trace_id": "source_trace_id",
    }.items():
        if key in updates:
            fields.append(f"{column} = {marker}")
            params.append(updates[key])
    for key, column in {"validation": "validation_json", "payload": "payload_json"}.items():
        if key in updates:
            fields.append(f"{column} = {marker}")
            params.append(_safe_json(updates[key]))
    params.append(case_id)
    with _lock:
        with _connect(settings) as conn:
            conn.execute(f"UPDATE curation_cases SET {', '.join(fields)} WHERE id = {marker}", params)
    return get_curation_case(settings, case_id)


def _refresh_case_status(settings: Settings, case_id: str) -> Optional[Dict[str, Any]]:
    case = get_curation_case(settings, case_id)
    if not case:
        return None
    drafts = case.get("drafts") or []
    validation = case.get("validation") or {}
    if case.get("status") == "resolved":
        return case
    if drafts and all(draft.get("status") in DOC_DONE_STATUSES for draft in drafts):
        next_status = "merged"
    elif any(draft.get("status") in {"patch_proposed", "merged"} for draft in drafts):
        next_status = "reviewing_patches"
    elif any(draft.get("status") in {"ai_generated", "draft_committed"} for draft in drafts):
        next_status = "drafting"
    elif drafts:
        next_status = "targets_selected"
    else:
        next_status = "new"
    if (validation.get("reindex") or {}).get("state") == "success":
        next_status = "reindexed"
    return _update_case(settings, case_id, {"status": next_status})


def refresh_curation_case_status(settings: Settings, case_id: str) -> Optional[Dict[str, Any]]:
    return _refresh_case_status(settings, case_id)


def _targets_from_trace(trace: Optional[Dict[str, Any]]) -> List[str]:
    targets: List[str] = []
    if not trace:
        return targets
    raw_sources: List[Any] = []
    raw_sources.extend(trace.get("selected_sources") or [])
    for item in trace.get("evidence") or []:
        if isinstance(item, dict):
            raw_sources.append(item.get("source"))
    for item in trace.get("candidates") or []:
        if isinstance(item, dict):
            raw_sources.append(item.get("source"))
    for source in raw_sources:
        repo_path = _canonical_repo_path_from_source(str(source or ""))
        if repo_path and repo_path not in targets:
            targets.append(repo_path)
    return targets[:8]


def _create_case_record(settings: Settings, payload: Dict[str, Any]) -> Dict[str, Any]:
    now = _now()
    case = {
        "id": str(uuid4()),
        "created_at": now,
        "updated_at": now,
        "status": payload.get("status") or "new",
        "title": (payload.get("title") or "Caso de curadoria").strip(),
        "instruction": (payload.get("instruction") or "").strip(),
        "source_kind": payload.get("source_kind") or "manual",
        "source_feedback_id": payload.get("source_feedback_id"),
        "source_trace_id": payload.get("source_trace_id"),
        "validation": payload.get("validation") or {},
        "payload": payload.get("payload") or {},
    }
    marker = _placeholder(settings)
    with _lock:
        with _connect(settings) as conn:
            conn.execute(
                f"""
                INSERT INTO curation_cases
                (id, created_at, updated_at, status, title, instruction, source_kind, source_feedback_id,
                 source_trace_id, validation_json, payload_json)
                VALUES ({",".join([marker] * 11)})
                """,
                (
                    case["id"],
                    case["created_at"],
                    case["updated_at"],
                    case["status"],
                    case["title"],
                    case["instruction"],
                    case["source_kind"],
                    case["source_feedback_id"],
                    case["source_trace_id"],
                    _safe_json(case["validation"]),
                    _safe_json(case["payload"]),
                ),
            )
    return get_curation_case(settings, case["id"]) or case


def add_targets_to_case(settings: Settings, case_id: str, target_paths: List[str]) -> Dict[str, Any]:
    case = get_curation_case(settings, case_id)
    if not case:
        raise AppError("curation_case_not_found", "Caso de curadoria não encontrado.", 404)
    existing = {draft.get("target_path") for draft in case.get("drafts") or []}
    created = []
    for raw_path in target_paths:
        target_path = _normalize_repo_path(raw_path)
        if target_path in existing:
            continue
        local_target = _local_path_for_repo_path(settings, target_path)
        if not local_target.exists():
            raise AppError("canonical_doc_not_found", "Markdown canônico não encontrado.", 404)
        title = f"{case['title']}: {PurePosixPath(target_path).name}"
        draft = create_draft(
            settings,
            {
                "title": title,
                "instruction": case.get("instruction") or "Curar documento canônico.",
                "case_id": case_id,
                "source_feedback_id": case.get("source_feedback_id"),
                "trace_id": case.get("source_trace_id"),
                "target_path": target_path,
                "canonical_targets": [target_path],
                "decision": "edit",
                "status": "review",
            },
        )
        created.append(draft)
        existing.add(target_path)
    _refresh_case_status(settings, case_id)
    return {"case": get_curation_case(settings, case_id), "created": created}


def create_curation_case(settings: Settings, payload: Dict[str, Any]) -> Dict[str, Any]:
    trace_id = payload.get("trace_id")
    trace = get_rag_trace(settings, trace_id) if trace_id else None
    targets = payload.get("target_paths") or _targets_from_trace(trace)
    case = _create_case_record(
        settings,
        {
            "title": payload.get("title") or "Curadoria manual",
            "instruction": payload.get("instruction") or "Atualizar a base de conhecimento com curadoria assistida.",
            "source_kind": payload.get("source_kind") or "manual",
            "source_feedback_id": payload.get("source_feedback_id"),
            "source_trace_id": trace_id,
            "payload": {"trace": trace, **(payload.get("payload") or {})},
        },
    )
    if targets:
        add_targets_to_case(settings, case["id"], targets)
    return get_curation_case(settings, case["id"]) or case


def create_case_from_feedback(settings: Settings, feedback_id: str) -> Dict[str, Any]:
    existing = active_case_for_feedback(settings, feedback_id)
    if existing:
        return {"case": existing, "created": False}
    feedback = get_rag_feedback(settings, feedback_id)
    if not feedback:
        raise AppError("rag_feedback_not_found", "Feedback RAG não encontrado.", 404)
    trace_id = feedback.get("trace_id")
    trace = get_rag_trace(settings, trace_id) if trace_id else None
    title_piece = feedback.get("reason") or feedback.get("rating") or "feedback"
    instruction = "\n\n".join(
        item
        for item in [
            f"Feedback: {feedback.get('rating')} / {feedback.get('reason')}",
            feedback.get("comment"),
            feedback.get("expected_answer"),
            f"Pergunta original: {trace.get('question')}" if trace else None,
        ]
        if item
    )
    case = create_curation_case(
        settings,
        {
            "title": f"Curadoria RAG: {title_piece}",
            "instruction": instruction or "Corrigir ou completar o contexto apontado pelo feedback.",
            "source_kind": "feedback",
            "source_feedback_id": feedback_id,
            "trace_id": trace_id,
            "payload": {"feedback": feedback},
        },
    )
    triage_rag_feedback(settings, feedback_id, "case_open", {"case_id": case["id"], "opened_at": _now()})
    return {"case": case, "created": True}


def ignore_case_draft(settings: Settings, case_id: str, draft_id: str) -> Dict[str, Any]:
    draft = get_draft(settings, draft_id)
    if not draft or draft.get("case_id") != case_id:
        raise AppError("draft_not_found", "Draft do caso não encontrado.", 404)
    update_draft(settings, draft_id, {"status": "ignored", "decision": "ignored", "completed_at": _now()})
    return {"case": _refresh_case_status(settings, case_id), "draft": get_draft(settings, draft_id, include_related=True)}


def record_case_reindex(settings: Settings, case_id: str, result: Dict[str, Any]) -> Optional[Dict[str, Any]]:
    case = get_curation_case(settings, case_id)
    if not case:
        return None
    validation = case.get("validation") or {}
    success = result.get("state") == "success" or result.get("state") == "idle" and not result.get("error")
    validation["reindex"] = {
        "state": "success" if success else "error",
        "at": _now(),
        "document_count": result.get("document_count"),
        "chunk_count": result.get("chunk_count"),
        "vector_backend": result.get("vector_backend"),
        "vector_chunks": result.get("vector_chunks"),
        "error": result.get("error"),
    }
    updated = _update_case(settings, case_id, {"validation": validation, "status": "reindexed" if success else "reindex_error"})
    if success and updated and updated.get("source_feedback_id"):
        triage_rag_feedback(settings, str(updated["source_feedback_id"]), "reindexed", {"case_id": case_id, "reindexed_at": _now()})
    return get_curation_case(settings, case_id)


def record_case_eval(settings: Settings, case_id: str, result: Dict[str, Any]) -> Optional[Dict[str, Any]]:
    case = get_curation_case(settings, case_id)
    if not case:
        return None
    validation = case.get("validation") or {}
    failed = int(result.get("failed") or 0)
    validation["eval"] = {
        "state": "passed" if failed == 0 else "failed",
        "at": _now(),
        "total": result.get("total"),
        "passed": result.get("passed"),
        "failed": result.get("failed"),
        "duration_ms": result.get("duration_ms"),
    }
    return _update_case(settings, case_id, {"validation": validation})


def resolve_curation_case(settings: Settings, case_id: str) -> Dict[str, Any]:
    case = get_curation_case(settings, case_id)
    if not case:
        raise AppError("curation_case_not_found", "Caso de curadoria não encontrado.", 404)
    drafts = case.get("drafts") or []
    if drafts and not all(draft.get("status") in DOC_DONE_STATUSES for draft in drafts):
        raise AppError("curation_case_incomplete", "Conclua ou ignore todos os documentos antes de resolver.", 422)
    if (case.get("validation") or {}).get("reindex", {}).get("state") != "success":
        raise AppError("curation_case_not_reindexed", "Reindexe e valide o caso antes de resolver.", 422)
    updated = _update_case(settings, case_id, {"status": "resolved"})
    if updated and updated.get("source_feedback_id"):
        triage_rag_feedback(settings, str(updated["source_feedback_id"]), "resolved", {"case_id": case_id, "resolved_at": _now()})
    return updated or case


def delete_draft(settings: Settings, draft_id: str) -> bool:
    marker = _placeholder(settings)
    with _lock:
        with _connect(settings) as conn:
            if _uses_postgres(settings):
                row = conn.execute(f"DELETE FROM knowledge_drafts WHERE id = {marker} RETURNING id", (draft_id,)).fetchone()
                conn.execute(f"DELETE FROM draft_attachments WHERE draft_id = {marker}", (draft_id,))
                conn.execute(f"DELETE FROM draft_runs WHERE draft_id = {marker}", (draft_id,))
                conn.execute(f"DELETE FROM draft_patches WHERE draft_id = {marker}", (draft_id,))
                return bool(row)
            cursor = conn.execute(f"DELETE FROM knowledge_drafts WHERE id = {marker}", (draft_id,))
            conn.execute(f"DELETE FROM draft_attachments WHERE draft_id = {marker}", (draft_id,))
            conn.execute(f"DELETE FROM draft_runs WHERE draft_id = {marker}", (draft_id,))
            conn.execute(f"DELETE FROM draft_patches WHERE draft_id = {marker}", (draft_id,))
            return cursor.rowcount > 0


def _count_attachments(settings: Settings, draft_id: str) -> int:
    marker = _placeholder(settings)
    with _lock:
        with _connect(settings) as conn:
            row = conn.execute(f"SELECT COUNT(*) AS total FROM draft_attachments WHERE draft_id = {marker}", (draft_id,)).fetchone()
    if isinstance(row, dict):
        return int(row.get("total") or 0)
    return int(row[0] if row else 0)


def _extract_json_text(data: bytes) -> str:
    parsed = json.loads(data.decode("utf-8"))
    return json.dumps(parsed, ensure_ascii=False, indent=2)


def _extract_pdf_text(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise AppError("pdf_parser_missing", "Parser PDF indisponivel. Instale pypdf no backend.", 503) from exc
    reader = PdfReader(io.BytesIO(data))
    return "\n\n".join((page.extract_text() or "").strip() for page in reader.pages if (page.extract_text() or "").strip())


def _extract_docx_text(data: bytes) -> str:
    try:
        import docx
    except ImportError as exc:
        raise AppError("docx_parser_missing", "Parser DOCX indisponivel. Instale python-docx no backend.", 503) from exc
    document = docx.Document(io.BytesIO(data))
    return "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())


def extract_attachment_text(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED_ATTACHMENT_SUFFIXES:
        raise AppError("unsupported_attachment", "Anexo precisa ser .md, .txt, .json, .pdf ou .docx.", 422)
    if suffix in {".md", ".txt"}:
        return data.decode("utf-8", errors="replace")
    if suffix == ".json":
        return _extract_json_text(data)
    if suffix == ".pdf":
        return _extract_pdf_text(data)
    if suffix == ".docx":
        return _extract_docx_text(data)
    return ""


def add_attachment(settings: Settings, draft_id: str, *, filename: str, content_type: str, data: bytes) -> Dict[str, Any]:
    if not get_draft(settings, draft_id):
        raise AppError("draft_not_found", "Draft nao encontrado.", 404)
    if _count_attachments(settings, draft_id) >= MAX_ATTACHMENTS_PER_DRAFT:
        raise AppError("too_many_attachments", "Limite de 5 anexos por draft atingido.", 422)
    max_bytes = MAX_ATTACHMENT_MB * 1024 * 1024
    if len(data) > max_bytes:
        raise AppError("attachment_too_large", "Anexo maior que 10 MB.", 422)
    extracted = extract_attachment_text(filename, data).strip()
    if not extracted:
        raise AppError("empty_attachment", "Nao consegui extrair texto deste anexo.", 422)
    attachment = {
        "id": str(uuid4()),
        "draft_id": draft_id,
        "created_at": _now(),
        "filename": Path(filename).name,
        "content_type": content_type,
        "size_bytes": len(data),
        "sha256": hashlib.sha256(data).hexdigest(),
        "extracted_text": extracted[:120000],
        "metadata": {"suffix": Path(filename).suffix.lower(), "truncated": len(extracted) > 120000},
    }
    marker = _placeholder(settings)
    with _lock:
        with _connect(settings) as conn:
            conn.execute(
                f"""
                INSERT INTO draft_attachments
                (id, draft_id, created_at, filename, content_type, size_bytes, sha256, extracted_text, metadata_json)
                VALUES ({",".join([marker] * 9)})
                """,
                (
                    attachment["id"],
                    attachment["draft_id"],
                    attachment["created_at"],
                    attachment["filename"],
                    attachment["content_type"],
                    attachment["size_bytes"],
                    attachment["sha256"],
                    attachment["extracted_text"],
                    _safe_json(attachment["metadata"]),
                ),
            )
    return attachment


def delete_attachment(settings: Settings, draft_id: str, attachment_id: str) -> bool:
    marker = _placeholder(settings)
    with _lock:
        with _connect(settings) as conn:
            if _uses_postgres(settings):
                row = conn.execute(
                    f"DELETE FROM draft_attachments WHERE id = {marker} AND draft_id = {marker} RETURNING id",
                    (attachment_id, draft_id),
                ).fetchone()
                return bool(row)
            cursor = conn.execute(
                f"DELETE FROM draft_attachments WHERE id = {marker} AND draft_id = {marker}",
                (attachment_id, draft_id),
            )
            return cursor.rowcount > 0


def _attachments_for_prompt(draft: Dict[str, Any]) -> str:
    blocks = []
    for item in draft.get("attachments") or []:
        text = str(item.get("extracted_text") or "")
        blocks.append(f"[attachment: {item.get('filename')}]\n{text[:8000]}")
    return "\n\n---\n\n".join(blocks)


def _canonical_target_contents(settings: Settings, targets: Iterable[str]) -> List[Dict[str, str]]:
    output = []
    for target in targets:
        repo_path = _normalize_repo_path(target)
        if repo_path.startswith(f"{DRAFT_ROOT}/"):
            continue
        local_path = _local_path_for_repo_path(settings, repo_path)
        if not local_path.exists():
            continue
        output.append({"path": repo_path, "content": local_path.read_text(encoding="utf-8")})
    return output[:3]


def list_canonical_documents(settings: Settings) -> List[Dict[str, Any]]:
    base = settings.resolved_knowledge_dir.resolve()
    docs: List[Dict[str, Any]] = []
    if not base.exists():
        return docs
    for path in sorted(base.rglob("*.md")):
        try:
            relative = path.resolve().relative_to(base).as_posix()
        except ValueError:
            continue
        if relative.startswith(("_drafts/", "_context/")):
            continue
        repo_path = f"knowledge/{relative}"
        text = path.read_text(encoding="utf-8")
        meta: Dict[str, Any] = {}
        if text.startswith("---"):
            parts = text.split("---", 2)
            if len(parts) >= 3:
                try:
                    parsed = yaml.safe_load(parts[1]) or {}
                    meta = parsed if isinstance(parsed, dict) else {}
                except yaml.YAMLError:
                    meta = {}
        if str(meta.get("visibility", "public")).strip().lower() != "public":
            continue
        docs.append(
            {
                "path": repo_path,
                "title": meta.get("title") or path.stem.replace("-", " ").title(),
                "category": meta.get("category") or relative.split("/")[0],
                "summary": meta.get("summary") or "",
                "tags": meta.get("tags") or [],
                "priority": meta.get("priority"),
            }
        )
    return docs


def _extract_json_object(text: str) -> Dict[str, Any]:
    cleaned = text.strip()
    if cleaned.startswith("```"):
        cleaned = re.sub(r"^```(?:json)?", "", cleaned, flags=re.IGNORECASE).strip()
        cleaned = re.sub(r"```$", "", cleaned).strip()
    try:
        value = json.loads(cleaned)
        return value if isinstance(value, dict) else {}
    except json.JSONDecodeError:
        start = cleaned.find("{")
        end = cleaned.rfind("}")
        if start >= 0 and end > start:
            try:
                value = json.loads(cleaned[start : end + 1])
                return value if isinstance(value, dict) else {}
            except json.JSONDecodeError:
                return {}
    return {}


def _fallback_agent_output(
    draft: Dict[str, Any],
    instruction: str,
    probe: Dict[str, Any],
    target_docs: List[Dict[str, str]],
    attachments_text: str,
) -> Dict[str, Any]:
    title = draft.get("title") or "Draft de curadoria"
    evidence = probe.get("evidence") or []
    frontmatter = _frontmatter_for_draft(title, "Draft gerado para revisar lacuna ou melhoria de conhecimento.")
    sources = "\n".join(f"- {item.get('source')}: {item.get('summary') or item.get('excerpt')}" for item in evidence[:6])
    body = f"""# {title}

## Objetivo

{instruction}

## Contexto recuperado

{sources or "Nenhuma evidência recuperada."}

## Anexos considerados

{attachments_text[:3000] if attachments_text else "Nenhum anexo informado."}

## Proposta de atualização

Revisar os documentos canônicos relacionados e incorporar apenas fatos confirmados. Este texto foi gerado como fallback estruturado e precisa de validação humana antes de qualquer merge.
"""
    proposed_markdown = _markdown_from_frontmatter(frontmatter, body)
    target_paths = [item["path"] for item in target_docs]
    return {
        "proposed_markdown": proposed_markdown,
        "target_sources": target_paths,
        "proposed_patches": [],
        "conflict_report": [],
        "evidence": evidence,
        "eval_case_suggestions": [
            {
                "question": instruction[:500],
                "expected_sources": [path.removeprefix("knowledge/") for path in target_paths[:3]],
                "forbidden_terms": [],
                "min_docs": 1,
            }
        ],
        "confidence": 0.45,
    }


def _agent_prompt(
    draft: Dict[str, Any],
    instruction: str,
    probe: Dict[str, Any],
    target_docs: List[Dict[str, str]],
    attachments_text: str,
) -> str:
    targets = "\n\n---\n\n".join(f"[target: {item['path']}]\n{item['content'][:10000]}" for item in target_docs)
    evidence = json.dumps(probe.get("evidence") or [], ensure_ascii=False, indent=2)[:14000]
    return f"""
Você é o Draft Studio Agent do portfólio de Gabriel.

Missão:
- Transformar feedback/instrução em uma proposta de curadoria RAG privada.
- Usar apenas contexto recuperado, documentos canônicos e anexos enviados.
- Não inventar fatos.
- Detectar conflitos em vez de decidir sozinho.
- Propor patch em documentos canônicos existentes, não criar documento público duplicado.

Instrução do Gabriel:
{instruction}

Draft atual:
{str(draft.get("draft_markdown") or "")[:12000]}

Evidências recuperadas pelo RAG:
{evidence}

Documentos canônicos alvo:
{targets or "Nenhum documento alvo encontrado."}

Anexos extraídos:
{attachments_text[:16000] if attachments_text else "Nenhum anexo enviado."}

Responda somente em JSON válido, sem markdown fora do JSON, neste formato:
{{
  "proposed_markdown": "Markdown draft completo com frontmatter visibility: draft",
  "target_sources": ["knowledge/...md"],
  "proposed_patches": [
    {{
      "target_path": "knowledge/...md",
      "proposed_content": "conteúdo completo do arquivo canônico após o patch",
      "rationale": "por que esse patch resolve a lacuna"
    }}
  ],
  "conflict_report": [
    {{
      "target_path": "knowledge/...md",
      "existing_excerpt": "trecho atual",
      "new_excerpt": "trecho novo",
      "decision_needed": "decisão humana necessária"
    }}
  ],
  "evidence": [],
  "eval_case_suggestions": [
    {{
      "question": "pergunta de regressão",
      "expected_sources": ["arquivo.md"],
      "forbidden_terms": [],
      "min_docs": 1
    }}
  ],
  "confidence": 0.0
}}
""".strip()


def _diff_text(path: str, before: str, after: str) -> str:
    return "\n".join(
        difflib.unified_diff(
            before.splitlines(),
            after.splitlines(),
            fromfile=f"a/{path}",
            tofile=f"b/{path}",
            lineterm="",
        )
    )


def _store_run(settings: Settings, draft_id: str, payload: Dict[str, Any]) -> Dict[str, Any]:
    run = {
        "id": str(uuid4()),
        "draft_id": draft_id,
        "created_at": _now(),
        "model": payload.get("model"),
        "instruction": payload.get("instruction"),
        "status": payload.get("status", "success"),
        "input": payload.get("input") or {},
        "output": payload.get("output") or {},
        "metrics": payload.get("metrics") or {},
        "error": payload.get("error"),
    }
    marker = _placeholder(settings)
    with _lock:
        with _connect(settings) as conn:
            conn.execute(
                f"""
                INSERT INTO draft_runs
                (id, draft_id, created_at, model, instruction, status, input_json, output_json, metrics_json, error)
                VALUES ({",".join([marker] * 10)})
                """,
                (
                    run["id"],
                    run["draft_id"],
                    run["created_at"],
                    run["model"],
                    run["instruction"],
                    run["status"],
                    _safe_json(run["input"]),
                    _safe_json(run["output"]),
                    _safe_json(run["metrics"]),
                    run["error"],
                ),
            )
    return run


def _store_patch(settings: Settings, draft_id: str, target_path: str, original: str, proposed: str, payload: Dict[str, Any]) -> Dict[str, Any]:
    patch = {
        "id": str(uuid4()),
        "draft_id": draft_id,
        "created_at": _now(),
        "status": "proposed",
        "target_path": _normalize_repo_path(target_path),
        "original_content": original,
        "proposed_content": proposed,
        "diff_text": _diff_text(target_path, original, proposed),
        "commit_sha": None,
        "payload": payload,
    }
    marker = _placeholder(settings)
    with _lock:
        with _connect(settings) as conn:
            conn.execute(
                f"""
                INSERT INTO draft_patches
                (id, draft_id, created_at, status, target_path, original_content, proposed_content, diff_text, commit_sha, payload_json)
                VALUES ({",".join([marker] * 10)})
                """,
                (
                    patch["id"],
                    patch["draft_id"],
                    patch["created_at"],
                    patch["status"],
                    patch["target_path"],
                    patch["original_content"],
                    patch["proposed_content"],
                    patch["diff_text"],
                    patch["commit_sha"],
                    _safe_json(patch["payload"]),
                ),
            )
    return patch


def get_patch(settings: Settings, draft_id: str, patch_id: str) -> Optional[Dict[str, Any]]:
    marker = _placeholder(settings)
    with _lock:
        with _connect(settings) as conn:
            row = conn.execute(
                f"SELECT * FROM draft_patches WHERE id = {marker} AND draft_id = {marker}",
                (patch_id, draft_id),
            ).fetchone()
    return _patch_from_row(row) if row else None


def update_patch_status(settings: Settings, draft_id: str, patch_id: str, status: str, commit_sha: Optional[str] = None) -> Optional[Dict[str, Any]]:
    marker = _placeholder(settings)
    with _lock:
        with _connect(settings) as conn:
            conn.execute(
                f"UPDATE draft_patches SET status = {marker}, commit_sha = {marker} WHERE id = {marker} AND draft_id = {marker}",
                (status, commit_sha, patch_id, draft_id),
            )
    return get_patch(settings, draft_id, patch_id)


def generate_draft_with_agent(
    settings: Settings,
    draft_id: str,
    *,
    instruction: str,
    target_paths: Optional[List[str]] = None,
) -> Dict[str, Any]:
    start = time.perf_counter()
    draft = get_draft(settings, draft_id, include_related=True)
    if not draft:
        raise AppError("draft_not_found", "Draft nao encontrado.", 404)
    clean_instruction = (instruction or draft.get("instruction") or "").strip()
    if not clean_instruction:
        raise AppError("empty_instruction", "Informe uma instrucao para o agente de draft.", 400)

    probe = build_rag_probe(settings, clean_instruction, None, limit=10)
    evidence = probe.get("evidence") or []
    inferred_targets = []
    for item in evidence:
        repo_path = _canonical_repo_path_from_source(str(item.get("source") or ""))
        if repo_path and repo_path not in inferred_targets:
            inferred_targets.append(repo_path)
    targets = [_normalize_repo_path(path) for path in (target_paths or []) if path] or inferred_targets[:3]
    target_docs = _canonical_target_contents(settings, targets)
    attachments_text = _attachments_for_prompt(draft)

    output: Dict[str, Any]
    error = None
    try:
        llm = _build_llm(settings)
        response = llm.invoke(
            [
                SystemMessage(content="Você é um agente de curadoria RAGOps. Responda somente JSON válido."),
                HumanMessage(content=_agent_prompt(draft, clean_instruction, probe, target_docs, attachments_text)),
            ]
        )
        parsed = _extract_json_object(_response_content_to_text(response.content))
        if not parsed:
            raise ValueError("LLM nao retornou JSON valido")
        output = parsed
    except Exception as exc:
        error = str(exc)[:500]
        output = _fallback_agent_output(draft, clean_instruction, probe, target_docs, attachments_text)

    proposed_markdown = sanitize_answer_text(str(output.get("proposed_markdown") or "")) or draft.get("draft_markdown") or ""
    target_sources = [
        _normalize_repo_path(path)
        for path in (output.get("target_sources") or targets)
        if isinstance(path, str) and path.strip()
    ]
    conflict_report = output.get("conflict_report") if isinstance(output.get("conflict_report"), list) else []
    eval_suggestions = output.get("eval_case_suggestions") if isinstance(output.get("eval_case_suggestions"), list) else []
    evidence_payload = output.get("evidence") if isinstance(output.get("evidence"), list) else evidence

    update_draft(
        settings,
        draft_id,
        {
            "status": "ai_generated" if not error else "review",
            "instruction": clean_instruction,
            "draft_markdown": proposed_markdown,
            "proposed_markdown": proposed_markdown,
            "canonical_targets": target_sources,
            "evidence": evidence_payload,
            "conflict_report": conflict_report,
            "eval_suggestions": eval_suggestions,
            "payload": {**(draft.get("payload") or {}), "agent_error": error, "last_probe": probe},
        },
    )

    patches = []
    target_by_path = {item["path"]: item["content"] for item in target_docs}
    for patch in output.get("proposed_patches") or []:
        if not isinstance(patch, dict):
            continue
        raw_target_path = str(patch.get("target_path") or "").strip()
        if not raw_target_path:
            continue
        target_path = _normalize_repo_path(raw_target_path)
        proposed_content = str(patch.get("proposed_content") or "")
        if not target_path or not proposed_content or target_path not in target_by_path:
            continue
        patches.append(
            _store_patch(
                settings,
                draft_id,
                target_path,
                target_by_path[target_path],
                proposed_content,
                {"rationale": patch.get("rationale"), "source": "agent"},
            )
        )

    run = _store_run(
        settings,
        draft_id,
        {
            "model": settings.openai_chat_model,
            "instruction": clean_instruction,
            "status": "fallback" if error else "success",
            "input": {"targets": targets, "attachments": [item.get("filename") for item in draft.get("attachments") or []]},
            "output": output,
            "metrics": {"total_ms": round((time.perf_counter() - start) * 1000, 2), "probe_ms": probe.get("took_ms")},
            "error": error,
        },
    )
    updated = get_draft(settings, draft_id, include_related=True)
    return {"draft": updated, "run": run, "probe": probe, "patches": patches, "fallback": bool(error)}


def propose_patch_from_draft(settings: Settings, draft_id: str, target_path: Optional[str] = None) -> Dict[str, Any]:
    draft = get_draft(settings, draft_id, include_related=True)
    if not draft:
        raise AppError("draft_not_found", "Draft nao encontrado.", 404)
    existing_open = [patch for patch in draft.get("patches") or [] if patch.get("status") == "proposed"]
    if existing_open and not target_path:
        return {"draft": draft, "patch": existing_open[0], "created": False}
    target = _normalize_repo_path(target_path or (draft.get("canonical_targets") or [None])[0] or "")
    if not target:
        raise AppError("missing_patch_target", "Nenhum documento canonico alvo foi definido.", 422)
    original = _local_path_for_repo_path(settings, target).read_text(encoding="utf-8")
    addition = f"""

## Complemento proposto pela curadoria

> Revisar antes de publicar. Origem: Draft Studio {draft_id}.

{(draft.get("proposed_markdown") or draft.get("draft_markdown") or "").strip()}
"""
    proposed = original.rstrip() + "\n" + addition.strip() + "\n"
    patch = _store_patch(settings, draft_id, target, original, proposed, {"source": "fallback_append"})
    update_draft(settings, draft_id, {"status": "patch_proposed"})
    return {"draft": get_draft(settings, draft_id, include_related=True), "patch": patch, "created": True}


def revert_draft_step(settings: Settings, draft_id: str, step: str) -> Optional[Dict[str, Any]]:
    draft = get_draft(settings, draft_id, include_related=True)
    if not draft:
        return None
    if step == "review":
        return update_draft(settings, draft_id, {"status": "review", "git_path": None, "git_commit_sha": None})
    if step == "discard_patch":
        marker = _placeholder(settings)
        with _lock:
            with _connect(settings) as conn:
                conn.execute(
                    f"UPDATE draft_patches SET status = {marker} WHERE draft_id = {marker} AND status = {marker}",
                    ("discarded", draft_id, "proposed"),
                )
        return update_draft(settings, draft_id, {"status": "review"})
    if step == "archive":
        return update_draft(settings, draft_id, {"status": "archived"})
    if step == "restore":
        return update_draft(settings, draft_id, {"status": "review"})
    raise AppError("invalid_revert_step", "Etapa de retorno invalida.", 422)
